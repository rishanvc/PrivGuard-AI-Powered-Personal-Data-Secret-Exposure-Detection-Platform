"""XLSX: one segment per non-empty cell, with sheet, cell reference and column header (FR-025, FR-029).

Read in openpyxl's streaming read-only mode (NFR-004), with data_only=False so
formulas come back AS WRITTEN: =HYPERLINK("https://user:secret@...") holds a
credential that the formula's displayed result hides. Hidden sheets are read
too - hiding a sheet does not remove its data (FR-099).
"""

from __future__ import annotations

import datetime

from openpyxl import load_workbook

from .base import ExtractionError, Limits, Report, Segment
from .ooxml import check_zip_safety


def _as_text(value) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e16:
        return str(int(value))  # 4.4e15 stored as a float still reads as digits
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def extract_xlsx(path: str, report: Report, limits: Limits):
    check_zip_safety(path, limits)

    # A file HANDLE, not a path: given a path, openpyxl refuses names without
    # an .xlsx extension - and stored uploads deliberately have none (D-016).
    with open(path, "rb") as handle:
        try:
            workbook = load_workbook(handle, read_only=True, data_only=False, keep_links=False)
        except Exception:
            raise ExtractionError("The workbook could not be opened. It may be damaged.") from None
        try:
            cells = 0
            for sheet in workbook.worksheets:
                hidden = sheet.sheet_state != "visible"
                report.count("sheets")
                if hidden:
                    report.count("hidden_sheets")
                    report.note("Hidden sheets were read too: hiding a sheet does not remove its data.")
                # A declared size like A1:XFD1048576 would make openpyxl pad rows
                # with millions of empty cells; ignore it and follow the real data.
                sheet.reset_dimensions()

                headers: dict[int, str] | None = None
                header_row = 0
                for row in sheet.iter_rows():
                    row_values = {}
                    row_number = 0
                    for cell in row:
                        # Gaps in a row come back as EmptyCell placeholders,
                        # which have no value - and no row/coordinate either.
                        if cell.value is None:
                            continue
                        text = _as_text(cell.value)
                        if not text.strip():
                            continue
                        cells += 1
                        if cells > limits.max_cells:
                            report.partial(f"Only the first {limits.max_cells:,} cells were read (the per-file cell limit).")
                            return
                        row_number = cell.row
                        row_values[cell.column] = text.strip()
                        location = {"sheet": sheet.title, "cell": cell.coordinate}
                        if hidden:
                            location["hidden"] = True
                        if headers and cell.row > header_row and headers.get(cell.column):
                            location["header"] = headers[cell.column]
                        yield Segment(text, location)
                    if headers is None and row_values:
                        headers, header_row = row_values, row_number  # first row with data
            report.set("cells", cells)
        finally:
            workbook.close()
