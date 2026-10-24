"""Build small, valid, synthetic Office and PDF files in memory.

Used by the engine tests and by samples/make_extraction_demo.py. Every value
these produce is synthetic - never real personal data (SRS 12.4).
"""

from __future__ import annotations

import io
import zipfile
from itertools import count
from xml.sax.saxutils import escape

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'

_revision_ids = count(1)


def zip_bytes(parts: dict[str, str | bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def _core_properties(creator: str | None) -> str:
    return (
        f'{XML_HEAD}<cp:coreProperties '
        'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f"<dc:creator>{escape(creator or '')}</dc:creator></cp:coreProperties>"
    )


# ------------------------------------------------------------------ Word

def run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


def deleted_run(text: str) -> str:
    """Text removed with Track Changes on - Word hides it, the file keeps it."""
    return (
        f'<w:del w:id="{next(_revision_ids)}" w:author="Synthetic Editor" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:delText xml:space="preserve">{escape(text)}</w:delText></w:r></w:del>'
    )


def field_runs(instruction: str, shown: str) -> str:
    """A field such as HYPERLINK: Word displays `shown`; the instruction stays in the file."""
    return (
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve"> {escape(instruction)} </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f"{run(shown)}"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )


def para(*runs: str) -> str:
    return f"<w:p>{''.join(runs)}</w:p>"


def tab_stop_para(before: str, after: str) -> str:
    """A paragraph that DEFINES a tab stop (in w:pPr) and uses one tab (in a run)."""
    return (
        '<w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="2880"/></w:tabs></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{escape(before)}</w:t><w:tab/>'
        f'<w:t xml:space="preserve">{escape(after)}</w:t></w:r></w:p>'
    )


def text_box_para(outer: str, inner: str) -> str:
    """A paragraph containing a text box, which contains a paragraph of its own."""
    return (
        f"<w:p>{run(outer)}<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        f"{para(run(inner))}</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )


def commented_para(text: str, comment_id: int = 0) -> str:
    return (
        f'<w:p><w:commentRangeStart w:id="{comment_id}"/>{run(text)}'
        f'<w:commentRangeEnd w:id="{comment_id}"/>'
        f'<w:r><w:commentReference w:id="{comment_id}"/></w:r></w:p>'
    )


def table(rows: list[list[str]]) -> str:
    columns = max(len(r) for r in rows)
    grid = "".join('<w:gridCol w:w="2400"/>' for _ in range(columns))
    body = "".join(
        "<w:tr>" + "".join(f"<w:tc>{para(run(cell))}</w:tc>" for cell in r) + "</w:tr>" for r in rows
    )
    return f'<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>{body}</w:tbl>'


def docx_bytes(
    body: str,
    *,
    header: str | None = None,
    footer: str | None = None,
    comments: list[tuple[str, str]] = (),
    creator: str | None = None,
    document_xml: str | None = None,
) -> bytes:
    """A Word document. `body` is WordprocessingML for inside <w:body>.

    `document_xml` replaces word/document.xml entirely (for malformed-input tests).
    """
    namespaces = f'xmlns:w="{W_NS}" xmlns:r="{R_NS}" xmlns:v="{V_NS}"'
    rels, overrides, sect = [], [], []
    parts: dict[str, str | bytes] = {}

    if header is not None:
        parts["word/header1.xml"] = f"{XML_HEAD}<w:hdr {namespaces}>{para(run(header))}</w:hdr>"
        rels.append(f'<Relationship Id="rIdHeader1" Type="{DOC_REL}/header" Target="header1.xml"/>')
        overrides.append(("/word/header1.xml", "wordprocessingml.header+xml"))
        sect.append('<w:headerReference w:type="default" r:id="rIdHeader1"/>')
    if footer is not None:
        parts["word/footer1.xml"] = f"{XML_HEAD}<w:ftr {namespaces}>{para(run(footer))}</w:ftr>"
        rels.append(f'<Relationship Id="rIdFooter1" Type="{DOC_REL}/footer" Target="footer1.xml"/>')
        overrides.append(("/word/footer1.xml", "wordprocessingml.footer+xml"))
        sect.append('<w:footerReference w:type="default" r:id="rIdFooter1"/>')
    if comments:
        items = "".join(
            f'<w:comment w:id="{i}" w:author="{escape(author)}" w:date="2026-01-01T00:00:00Z" w:initials="SR">'
            f"{para(run(text))}</w:comment>"
            for i, (author, text) in enumerate(comments)
        )
        parts["word/comments.xml"] = f"{XML_HEAD}<w:comments {namespaces}>{items}</w:comments>"
        rels.append(f'<Relationship Id="rIdComments" Type="{DOC_REL}/comments" Target="comments.xml"/>')
        overrides.append(("/word/comments.xml", "wordprocessingml.comments+xml"))

    parts["word/document.xml"] = document_xml or (
        f"{XML_HEAD}<w:document {namespaces}><w:body>{body}<w:sectPr>{''.join(sect)}</w:sectPr></w:body></w:document>"
    )
    parts["word/_rels/document.xml.rels"] = f'{XML_HEAD}<Relationships xmlns="{REL_NS}">{"".join(rels)}</Relationships>'

    package_rels = [f'<Relationship Id="rIdDoc" Type="{DOC_REL}/officeDocument" Target="word/document.xml"/>']
    if creator is not None:
        parts["docProps/core.xml"] = _core_properties(creator)
        package_rels.append(
            '<Relationship Id="rIdCore" Type="http://schemas.openxmlformats.org/package/2006/'
            'relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        )
        overrides.append(("/docProps/core.xml", None))
    parts["_rels/.rels"] = f'{XML_HEAD}<Relationships xmlns="{REL_NS}">{"".join(package_rels)}</Relationships>'

    def content_type(kind):
        if kind is None:
            return "application/vnd.openxmlformats-package.core-properties+xml"
        return f"application/vnd.openxmlformats-officedocument.{kind}"

    override_xml = "".join(
        f'<Override PartName="{name}" ContentType="{content_type(kind)}"/>' for name, kind in overrides
    )
    parts["[Content_Types].xml"] = (
        f'{XML_HEAD}<Types xmlns="{CT_NS}">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        f"{override_xml}</Types>"
    )
    return zip_bytes(parts)


# ------------------------------------------------------------------ PowerPoint

def _drawing_paragraphs(lines: list[str]) -> str:
    return "".join(
        "<a:p>" + "<a:br/>".join(f"<a:r><a:t>{escape(piece)}</a:t></a:r>" for piece in line.split("\n")) + "</a:p>"
        for line in lines
    )


def _shape_tree(lines: list[str]) -> str:
    return (
        "<p:cSld><p:spTree><p:sp><p:txBody><a:bodyPr/>"
        f"{_drawing_paragraphs(lines)}</p:txBody></p:sp></p:spTree></p:cSld>"
    )


def pptx_bytes(
    slides: dict[int, list[str]],
    *,
    order: list[int] | None = None,
    notes: dict[int, str] | None = None,
    creator: str | None = None,
) -> bytes:
    """A presentation. `slides` maps FILE number (slideN.xml) to its text lines;
    `order` lists file numbers in PRESENTATION order (defaults to ascending).

    Only the parts PrivGuard reads are produced - no masters, layouts or theme -
    so PowerPoint itself may refuse to open the result.
    """
    order = order or sorted(slides)
    notes = notes or {}
    ns = f'xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}"'
    parts: dict[str, str | bytes] = {}

    slide_ids = "".join(f'<p:sldId id="{256 + i}" r:id="rIdSlide{n}"/>' for i, n in enumerate(order))
    parts["ppt/presentation.xml"] = f"{XML_HEAD}<p:presentation {ns}><p:sldIdLst>{slide_ids}</p:sldIdLst></p:presentation>"
    parts["ppt/_rels/presentation.xml.rels"] = (
        f'{XML_HEAD}<Relationships xmlns="{REL_NS}">'
        + "".join(f'<Relationship Id="rIdSlide{n}" Type="{DOC_REL}/slide" Target="slides/slide{n}.xml"/>' for n in slides)
        + "</Relationships>"
    )
    for n, lines in slides.items():
        parts[f"ppt/slides/slide{n}.xml"] = f"{XML_HEAD}<p:sld {ns}>{_shape_tree(lines)}</p:sld>"
        if n in notes:
            parts[f"ppt/notesSlides/notesSlide{n}.xml"] = f"{XML_HEAD}<p:notes {ns}>{_shape_tree([notes[n]])}</p:notes>"
            parts[f"ppt/slides/_rels/slide{n}.xml.rels"] = (
                f'{XML_HEAD}<Relationships xmlns="{REL_NS}">'
                f'<Relationship Id="rIdNotes" Type="{DOC_REL}/notesSlide" Target="../notesSlides/notesSlide{n}.xml"/>'
                "</Relationships>"
            )
    if creator is not None:
        parts["docProps/core.xml"] = _core_properties(creator)
    parts["[Content_Types].xml"] = (
        f'{XML_HEAD}<Types xmlns="{CT_NS}">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/presentation.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
        "</Types>"
    )
    return zip_bytes(parts)


# ------------------------------------------------------------------ Excel

def xlsx_bytes(sheets: dict[str, list[list]], *, hidden: tuple[str, ...] = ()) -> bytes:
    """A workbook written by openpyxl. Strings starting with "=" become formulas."""
    from openpyxl import Workbook

    workbook = Workbook()
    for index, (title, rows) in enumerate(sheets.items()):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        for row in rows:
            sheet.append(row)
        if title in hidden:
            sheet.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ PDF

def _pdf_string(text: str) -> str:
    return "(" + text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def pdf_bytes(pages: list[str], *, info: dict[str, str] | None = None) -> bytes:
    """A valid PDF with a correct cross-reference table.

    An empty string makes a page with a drawn line but no text - standing in
    for a scanned page. Lines within a page are separated by "\\n".
    """
    objects: list[bytes] = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        if text:
            lines = " T* ".join(f"{_pdf_string(line)} Tj" for line in text.split("\n"))
            stream = f"BT /F1 12 Tf 14 TL 72 720 Td {lines} ET".encode("latin-1")
        else:
            stream = b"0.5 w 72 700 m 540 700 l S"  # graphics only: no text
        page_number = len(objects) + 1
        content_number = page_number + 1
        kids.append(page_number)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_number} 0 R "
            "/Resources << /Font << /F1 3 0 R >> >> >>".encode("latin-1")
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = ("<< /Type /Pages /Kids [" + " ".join(f"{k} 0 R" for k in kids) + f"] /Count {len(kids)} >>").encode()

    info_ref = ""
    if info:
        objects.append(("<< " + " ".join(f"/{k} {_pdf_string(v)}" for k, v in info.items()) + " >>").encode("latin-1"))
        info_ref = f" /Info {len(objects)} 0 R"

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R{info_ref} >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


# ---------------------------------------------------------------- images and scans (Cycle 8)

def _render(text: str, dpi: int = 150):
    """Draw text into a one-page PDF and render it to a grayscale bitmap."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(pdf_bytes([text]))
    page = document[0]
    bitmap = page.render(scale=dpi / 72, grayscale=True)
    page.close()
    return document, bitmap


def png_bytes(text: str, *, dpi: int = 150) -> bytes:
    """A real PNG image showing `text` - built with zlib alone, no image library.

    A PNG is a signature, then chunks: IHDR (size and pixel format), IDAT (the
    rows, each prefixed with a filter byte, compressed with zlib) and IEND.
    """
    import struct
    import zlib

    document, bitmap = _render(text, dpi)
    try:
        raw = bytes(bitmap.buffer)
        rows = b"".join(b"\x00" + raw[y * bitmap.stride: y * bitmap.stride + bitmap.width] for y in range(bitmap.height))
        width, height = bitmap.width, bitmap.height
    finally:
        bitmap.close()
        document.close()

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    header = struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)  # 8-bit grayscale, no interlace
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows, 9)) + chunk(b"IEND", b"")


def png_header_only(width: int, height: int) -> bytes:
    """The first bytes of a PNG claiming a size - enough to test the size check, with no pixels."""
    import struct

    return b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)


def jpeg_header_only(width: int, height: int) -> bytes:
    """A JPEG's opening markers up to its frame header: enough to read its size, with no image data."""
    import struct

    app0 = b"\xff\xe0" + struct.pack(">H", 16) + b"JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    sof0 = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, height, width, 1) + b"\x01\x11\x00"
    return b"\xff\xd8" + app0 + sof0


def scanned_pdf_bytes(pages: list[str], *, dpi: int = 150) -> bytes:
    """A PDF whose pages are only PICTURES of text - no text layer, as a scanner makes."""
    import pypdfium2 as pdfium

    scanned = pdfium.PdfDocument.new()
    try:
        for text in pages:
            document, bitmap = _render(text, dpi)
            try:
                page = scanned.new_page(612, 792)
                image = pdfium.PdfImage.new(scanned)
                image.set_bitmap(bitmap)
                image.set_matrix(pdfium.PdfMatrix().scale(612, 792))  # stretch the picture over the page
                page.insert_obj(image)
                page.gen_content()
                page.close()
            finally:
                bitmap.close()
                document.close()
        out = io.BytesIO()
        scanned.save(out)
        return out.getvalue()
    finally:
        scanned.close()
